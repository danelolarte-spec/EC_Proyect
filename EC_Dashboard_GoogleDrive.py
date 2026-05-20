#!/usr/bin/env python3
"""
EC_Dashboard_GoogleDrive.py

Envia el .xlsx generado por el dashboard EC Transportes por correo
usando SMTP de Gmail con contrasena de aplicacion.

Uso:
    python EC_Dashboard_GoogleDrive.py archivo.xlsx

Requisitos:
    pip install openpyxl

Configuracion previa (una sola vez):
    1. En tu cuenta Google: activa la verificacion en 2 pasos
       (myaccount.google.com -> Seguridad).
    2. Genera una "Contrasena de aplicacion" en
       myaccount.google.com/apppasswords (elige "Correo" / "Otro").
       Obtienes 16 caracteres tipo: abcd efgh ijkl mnop
    3. Edita GMAIL_USER y GMAIL_APP_PASSWORD abajo, o exportalos
       como variables de entorno antes de ejecutar:
           export GMAIL_USER="tu_correo@gmail.com"
           export GMAIL_APP_PASSWORD="abcd efgh ijkl mnop"
"""
import sys
import os
import smtplib
import ssl
from email.message import EmailMessage
from pathlib import Path
from datetime import datetime

try:
    from openpyxl import load_workbook
except ImportError:
    print("Falta openpyxl. Instala con:  pip install openpyxl")
    try:
        input("\nPresiona Enter para cerrar...")
    except EOFError:
        pass
    sys.exit(1)

# =============================================================
# CONFIGURACION  (edita estos valores)
# =============================================================
GMAIL_USER = os.environ.get("GMAIL_USER", "tu_correo@gmail.com")
GMAIL_APP_PASSWORD = os.environ.get("GMAIL_APP_PASSWORD", "xxxx xxxx xxxx xxxx")

# Destinatarios fijos
DESTINATARIOS = [
    "gerenciademarca@ectransportes.com",
]

# Con copia (opcional, dejar lista vacia [] si no se usa)
CC = []

EMPRESA = "EC Transportes"
SMTP_HOST = "smtp.gmail.com"
SMTP_PORT = 465
# =============================================================


MONTHS_ES = ['Enero', 'Febrero', 'Marzo', 'Abril', 'Mayo', 'Junio',
             'Julio', 'Agosto', 'Septiembre', 'Octubre', 'Noviembre', 'Diciembre']


def fmt_cop(v):
    return "$" + f"{int(round(v)):,}".replace(",", ".")


def fmt_pct(v):
    return f"{v * 100:.1f}%"


def numv(v):
    try:
        return float(v) if v is not None else 0.0
    except (TypeError, ValueError):
        return 0.0


def compute_kpis(xlsx_path):
    wb = load_workbook(xlsx_path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        return None
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    data = [dict(zip(headers, r)) for r in rows[1:]]

    # Periodo (mes / anio) a partir de la primera fecha valida
    mes = "-"
    for r in data:
        d = r.get('FECHA DEL SERVICIO')
        if isinstance(d, datetime):
            mes = f"{MONTHS_ES[d.month - 1]} {d.year}"
            break

    active = [r for r in data if r.get('ESTADO') in ('DIALIQUIDADO', 'LIQUIDADO')]
    cancelled = [r for r in data if r.get('ESTADO') == 'CANCELADO']

    total_active = len(active)
    total_cancelled = len(cancelled)
    total_svc = len(data)
    pct_cancel = (total_cancelled / total_svc) if total_svc else 0.0
    total_factura = sum(numv(r.get('VALOR')) for r in active)
    ticket = (total_factura / total_active) if total_active else 0.0

    clientes = set()
    conductores = set()
    cliente_fact = {}
    conductor_fact = {}
    for r in active:
        c = r.get('CLIENTE')
        if c:
            clientes.add(c)
            cliente_fact[c] = cliente_fact.get(c, 0.0) + numv(r.get('VALOR'))
        d = r.get('CONDUCTOR')
        if d:
            conductores.add(d)
            conductor_fact[d] = conductor_fact.get(d, 0.0) + numv(r.get('VALOR'))

    top_cli = sorted(cliente_fact.items(), key=lambda x: -x[1])[:5]
    top_cond = sorted(conductor_fact.items(), key=lambda x: -x[1])[:5]

    return {
        'mes': mes,
        'total_active': total_active,
        'total_cancelled': total_cancelled,
        'pct_cancel': pct_cancel,
        'total_factura': total_factura,
        'ticket': ticket,
        'clientes_n': len(clientes),
        'conductores_n': len(conductores),
        'top_cli': top_cli,
        'top_cond': top_cond,
    }


def build_html(k):
    rows_cli = "".join(
        f"<tr><td style='padding:6px;border-bottom:1px solid #eee'>{i+1}. {n}</td>"
        f"<td style='padding:6px;border-bottom:1px solid #eee;text-align:right;color:#1A7A4A'>{fmt_cop(v)}</td></tr>"
        for i, (n, v) in enumerate(k['top_cli'])
    )
    rows_cond = "".join(
        f"<tr><td style='padding:6px;border-bottom:1px solid #eee'>{i+1}. {n}</td>"
        f"<td style='padding:6px;border-bottom:1px solid #eee;text-align:right;color:#1A7A4A'>{fmt_cop(v)}</td></tr>"
        for i, (n, v) in enumerate(k['top_cond'])
    )
    return f"""<!doctype html>
<html><body style="font-family:Arial,sans-serif;color:#1B2A4A;max-width:680px;margin:0 auto;padding:24px;background:#fff">
  <div style="border-bottom:3px solid #C9A84C;padding-bottom:14px;margin-bottom:22px">
    <h2 style="color:#C9A84C;margin:0;letter-spacing:2px">EC TRANSPORTES</h2>
    <p style="margin:4px 0 0;color:#666;font-size:13px">Dashboard Ejecutivo &mdash; {k['mes']}</p>
  </div>

  <h3 style="color:#1B2A4A;font-size:15px">Indicadores Clave</h3>
  <table cellpadding="0" cellspacing="0" style="border-collapse:collapse;width:100%;font-size:14px">
    <tr><td style="padding:8px;border-bottom:1px solid #eee"><strong>Servicios Liquidados</strong></td>
        <td style="padding:8px;border-bottom:1px solid #eee;text-align:right;color:#1A7A4A;font-weight:bold">{k['total_active']:,}</td></tr>
    <tr><td style="padding:8px;border-bottom:1px solid #eee"><strong>Cancelados</strong></td>
        <td style="padding:8px;border-bottom:1px solid #eee;text-align:right;color:#C0392B">{k['total_cancelled']:,} ({fmt_pct(k['pct_cancel'])})</td></tr>
    <tr><td style="padding:8px;border-bottom:1px solid #eee"><strong>Facturacion Total</strong></td>
        <td style="padding:8px;border-bottom:1px solid #eee;text-align:right;color:#C9A84C;font-weight:bold;font-size:16px">{fmt_cop(k['total_factura'])}</td></tr>
    <tr><td style="padding:8px;border-bottom:1px solid #eee"><strong>Ticket Promedio</strong></td>
        <td style="padding:8px;border-bottom:1px solid #eee;text-align:right">{fmt_cop(k['ticket'])}</td></tr>
    <tr><td style="padding:8px;border-bottom:1px solid #eee"><strong>Conductores Activos</strong></td>
        <td style="padding:8px;border-bottom:1px solid #eee;text-align:right">{k['conductores_n']}</td></tr>
    <tr><td style="padding:8px"><strong>Clientes Distintos</strong></td>
        <td style="padding:8px;text-align:right">{k['clientes_n']}</td></tr>
  </table>

  <h3 style="color:#1B2A4A;margin-top:26px;font-size:15px">Top 5 Clientes</h3>
  <table cellpadding="0" cellspacing="0" style="border-collapse:collapse;width:100%;font-size:13px">{rows_cli}</table>

  <h3 style="color:#1B2A4A;margin-top:26px;font-size:15px">Top 5 Conductores</h3>
  <table cellpadding="0" cellspacing="0" style="border-collapse:collapse;width:100%;font-size:13px">{rows_cond}</table>

  <p style="margin-top:26px;font-size:12px;color:#666;background:#f7f7f7;padding:12px;border-radius:6px">
    Adjunto: dashboard completo en Excel con todas las hojas de detalle
    (clientes, conductores, vehiculos, productos, hoteles, tendencia diaria,
    horas pico, cancelaciones).
  </p>

  <hr style="border:none;border-top:1px solid #eee;margin:24px 0">
  <p style="font-size:11px;color:#999">Reporte generado automaticamente &mdash; {EMPRESA}</p>
</body></html>
"""


def send_mail(xlsx_path, kpis):
    msg = EmailMessage()
    msg['Subject'] = f"Dashboard Ejecutivo {EMPRESA} - {kpis['mes']}"
    msg['From'] = GMAIL_USER
    msg['To'] = ", ".join(DESTINATARIOS)
    if CC:
        msg['Cc'] = ", ".join(CC)

    msg.set_content(
        f"Dashboard {EMPRESA} - {kpis['mes']}\n\n"
        f"Servicios liquidados: {kpis['total_active']}\n"
        f"Facturacion total: {fmt_cop(kpis['total_factura'])}\n"
        f"Cancelados: {kpis['total_cancelled']} ({fmt_pct(kpis['pct_cancel'])})\n\n"
        f"Ver el correo en formato HTML para mas detalles. Archivo adjunto."
    )
    msg.add_alternative(build_html(kpis), subtype='html')

    with open(xlsx_path, 'rb') as f:
        msg.add_attachment(
            f.read(),
            maintype='application',
            subtype='vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            filename=Path(xlsx_path).name,
        )

    recipients = DESTINATARIOS + CC
    ctx = ssl.create_default_context()
    with smtplib.SMTP_SSL(SMTP_HOST, SMTP_PORT, context=ctx) as smtp:
        smtp.login(GMAIL_USER, GMAIL_APP_PASSWORD)
        smtp.send_message(msg, from_addr=GMAIL_USER, to_addrs=recipients)


def pick_file_dialog():
    """Abre un dialogo de Windows/Mac/Linux para escoger el .xlsx."""
    try:
        import tkinter as tk
        from tkinter import filedialog
        root = tk.Tk()
        root.withdraw()
        root.attributes('-topmost', True)
        path = filedialog.askopenfilename(
            title="Selecciona el archivo .xlsx del dashboard",
            filetypes=[("Excel", "*.xlsx *.xls"), ("Todos", "*.*")],
        )
        root.destroy()
        return path
    except Exception:
        return None


def pause_exit(code=0):
    """Espera Enter antes de cerrar (para que no se cierre la ventana en Windows)."""
    try:
        input("\nPresiona Enter para cerrar...")
    except EOFError:
        pass
    sys.exit(code)


def main():
    print("=" * 60)
    print(" EC TRANSPORTES - Envio de Dashboard por correo")
    print("=" * 60)

    # 1) Obtener ruta del xlsx (argumento o dialogo)
    if len(sys.argv) >= 2:
        xlsx_path = sys.argv[1]
    else:
        print("\nNo pasaste un archivo. Abriendo selector...")
        xlsx_path = pick_file_dialog()
        if not xlsx_path:
            xlsx_path = input("Pega aqui la ruta del .xlsx: ").strip().strip('"')

    if not xlsx_path or not Path(xlsx_path).exists():
        print(f"\nERROR: no existe el archivo: {xlsx_path}")
        pause_exit(1)

    # 2) Validar credenciales
    if GMAIL_APP_PASSWORD == "xxxx xxxx xxxx xxxx" or GMAIL_USER == "tu_correo@gmail.com":
        print("\nERROR: faltan credenciales.")
        print("Abre este archivo .py con el Bloc de notas y edita:")
        print("   GMAIL_USER         = tu correo de Gmail")
        print("   GMAIL_APP_PASSWORD = la contrasena de aplicacion (16 chars)")
        print("\nPara generarla: myaccount.google.com/apppasswords")
        print("(Necesitas verificacion en 2 pasos activada).")
        pause_exit(1)

    # 3) Procesar y enviar
    print(f"\n[1/3] Leyendo {Path(xlsx_path).name} ...")
    try:
        kpis = compute_kpis(xlsx_path)
    except Exception as e:
        print(f"ERROR al leer el archivo: {e}")
        pause_exit(1)

    if not kpis:
        print("ERROR: el archivo esta vacio o no tiene filas de datos.")
        pause_exit(1)

    print(f"[2/3] Periodo: {kpis['mes']} | "
          f"{kpis['total_active']} servicios | {fmt_cop(kpis['total_factura'])}")

    print(f"[3/3] Enviando a: {', '.join(DESTINATARIOS)} ...")
    try:
        send_mail(xlsx_path, kpis)
    except smtplib.SMTPAuthenticationError:
        print("\nERROR: autenticacion fallida. Verifica:")
        print("  - Tener verificacion en 2 pasos activada en tu cuenta Google.")
        print("  - Usar contrasena de APLICACION (no la contrasena normal).")
        print("  - GMAIL_USER y GMAIL_APP_PASSWORD correctos.")
        pause_exit(1)
    except Exception as e:
        print(f"\nERROR al enviar: {e}")
        pause_exit(1)

    print("\nOK: correo enviado correctamente.")
    pause_exit(0)


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print("\nCancelado por el usuario.")
        pause_exit(1)
    except Exception as e:
        print(f"\nERROR inesperado: {e}")
        import traceback
        traceback.print_exc()
        pause_exit(1)
